from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fractions import Fraction as F
import json

src=Document(r'C:\Users\Dell\Downloads\Bai_thuc_hanh_EVM_45_phut_De_bai_sinh_vien.docx')
actual=[1,1,1,1,1,F(1,2),F(1,3),F(2,3),F(3,5),F(4,6),F(2,4),F(2,5),F(4,7),F(1,4),F(2,6),1,1,1,F(1,2),0]
cost=[600,720,500,780,300,550,500,1300,750,1500,600,1100,1000,500,700,600,950,1450,450,0]
data=[]; ends={}
for i,row in enumerate(src.tables[0].rows[1:]):
    id,name,pred,dur,rate=[c.text for c in row.cells]; dur=int(dur); rate=int(rate)
    start=ends.get(pred,0)+1; end=start+dur-1; ends[id]=end
    planned=F(min(dur,max(0,16-start)),dur); budget=dur*rate; ev=budget*actual[i]; pv=budget*planned; ac=cost[i]
    data.append(dict(id=id,name=name,pred=pred,dur=dur,rate=rate,start=start,end=end,p=planned,a=actual[i],b=budget,pv=pv,ev=ev,ac=ac,cv=ev-ac,sv=ev-pv,cpi=ev/ac if ac else None,spi=ev/pv if pv else None))
tot={k:sum(r[k] for r in data) for k in ['b','pv','ev','ac','cv','sv']}; tot['cpi']=tot['ev']/tot['ac']; tot['spi']=tot['ev']/tot['pv']
print({k:float(v) for k,v in tot.items()})
def num(v,n=2):
    if v is None:return 'KXĐ'
    s=f'{float(v):,.{n}f}'; return s.replace(',','@').replace('.',',').replace('@','.')
def money(v):return num(v,0) if float(v).is_integer() else num(v)
def pct(v):return num(v*100,0 if float(v*100).is_integer() else 2)+'%'
d=Document(); sec=d.sections[0]; sec.page_width=Cm(29.7); sec.page_height=Cm(21); sec.top_margin=sec.bottom_margin=Cm(1.4); sec.left_margin=sec.right_margin=Cm(1.5)
for sn in ['Normal','Title','Heading 1','Heading 2']:
    st=d.styles[sn]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0); st.font.size=Pt(10 if sn=='Normal' else 19 if sn=='Title' else 14 if sn=='Heading 1' else 11)
    st.paragraph_format.space_after=Pt(6)
def p(s):return d.add_paragraph(s)
def h(s):d.add_heading(s,1)
def table(headers,rows,widths=None,size=9):
    t=d.add_table(rows=1,cols=len(headers)); t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,s in zip(t.rows[0].cells,headers):c.text=str(s)
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=str(s)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trpr.append(cant)
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            if widths:c.width=Cm(widths[ci])
            tcpr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF3' if ri==0 else 'FFFFFF');tcpr.append(sh)
            for pa in c.paragraphs:
                pa.alignment=WD_ALIGN_PARAGRAPH.CENTER;pa.paragraph_format.space_after=Pt(3);pa.paragraph_format.space_before=Pt(3)
                for run in pa.runs:run.font.size=Pt(size);run.bold=ri==0
    return t
d.add_paragraph('Lời giải bài thực hành EVM', 'Title')
p('Dự án xây dựng Website thương mại điện tử • Báo cáo tại cuối ngày 15')
p('Kết luận: dự án chậm tiến độ và vượt chi phí so với giá trị công việc đã hoàn thành. SPI = '+num(tot['spi'])+'; CPI = '+num(tot['cpi'])+'. Lịch cơ sở kết thúc ngày 20.')
h('1 Lịch cơ sở và biểu đồ Gantt')
p('Quy ước: ngày bắt đầu và kết thúc đều được tính vào thời lượng. ES = EF tiền nhiệm + 1; EF = ES + Duration − 1. Ô xanh là thời gian kế hoạch; đường đỏ sau ngày 15 là thời điểm báo cáo.')
t=table(['Task','Trước','BĐ','KT']+[str(i) for i in range(1,21)],[[r['id'],r['pred'],r['start'],r['end']]+['' for _ in range(20)] for r in data],[1.1,1.1,.9,.9]+[1.13]*20,8)
for ri,r in enumerate(data,1):
    for day in range(r['start'],r['end']+1):
        sh=t.cell(ri,day+3)._tc.get_or_add_tcPr().find(qn('w:shd'));sh.set(qn('w:fill'),'648CAF')
for row in t.rows:
    edge=row.cells[18]._tc.get_or_add_tcPr().find(qn('w:tcBorders')).find(qn('w:right'));edge.set(qn('w:color'),'C00000');edge.set(qn('w:sz'),'18')
p('Đường găng theo quan hệ trong đề: A → B → C → D → E → P → Q → R → S → T = 20 ngày. Các nhánh còn lại kết thúc muộn nhất ngày 17.')
d.add_page_break();h('2 Thống kê tiến độ tại cuối ngày 15')
p('Giả định ngân sách và tiến độ kế hoạch phân bổ đều theo ngày. Planned % = min(Duration; max(0; 15 − ES + 1)) / Duration. Các phân số thực tế là tỷ lệ hoàn thành, không phải thời lượng kế hoạch.')
table(['ID','Công việc','BĐ','KT','Budget','Planned','Actual','AC','Trạng thái'],[[r['id'],r['name'],r['start'],r['end'],money(r['b']),pct(r['p']),pct(r['a']),money(r['ac']),'Chưa đến lịch' if r['p']==0 and r['a']==0 else 'Trễ' if r['a']<r['p'] else 'Đúng tiến độ' if r['a']==r['p'] else 'Vượt tiến độ'] for r in data],[.8,7.4,.9,.9,2.1,2.1,2.1,2.1,3.2],9)
p('10 công việc trễ: F, G, H, I, J, K, L, M, N, O. 9 công việc đúng tiến độ: A, B, C, D, E, P, Q, R, S. T chưa đến lịch. Riêng S đạt 50% đúng kế hoạch; không được kết luận trễ chỉ vì chưa hoàn thành.')
d.add_page_break();h('3 Bảng tính Earned Value Management')
p('Đơn vị chi phí giữ nguyên như đề. Budget = Duration × Cost/Day; PV = Budget × Planned %; EV = Budget × Actual %. CV = EV − AC; CPI = EV/AC; SV = EV − PV; SPI = EV/PV.')
keys=['id','b','pv','ev','ac','cv','cpi','sv','spi']
rows=[[r['id']]+[num(r[k]) if k in ['cpi','spi'] else money(r[k]) for k in keys[1:]] for r in data]
rows.append(['TỔNG']+[num(tot[k]) if k in ['cpi','spi'] else money(tot[k]) for k in keys[1:]])
table(['ID','Budget','PV','EV','AC','CV','CPI','SV','SPI'],rows,[1.1,3.2,3.2,3.2,3.1,3.1,2,3.1,2],9)
p('KXĐ = không xác định: task T có AC = PV = 0 nên CPI và SPI là 0/0. Tổng được tính từ số liệu chưa làm tròn; CPI và SPI toàn dự án là tỷ số của các tổng, không phải trung bình chỉ số từng task.')
p('Ví dụ O: kế hoạch ngày 14–17, nên Planned = 2/4 = 50%; Budget = 4 × 450 = 1.800; PV = 900; EV = 1.800 × 2/6 = 600; CV = −100; SV = −300. M có EV = 1.500 × 4/7 = 857,142857…')
d.add_page_break();h('4 Nhận định và quyết định của Project Manager')
d.add_heading('Đánh giá tiến độ và chi phí',2)
p('SV = '+money(tot['sv'])+' < 0 và SPI = '+num(tot['spi'])+' < 1: giá trị công việc hoàn thành chỉ bằng khoảng '+pct(tot['spi'])+' giá trị dự kiến tại ngày 15. Dự án chậm theo EVM; SV có đơn vị chi phí, không phải số ngày trễ. Chưa đủ dữ liệu để khẳng định ngày hoàn thành thực tế.')
p('CV = '+money(tot['cv'])+' < 0 và CPI = '+num(tot['cpi'])+' < 1: đã chi '+money(tot['ac'])+' để tạo ra giá trị '+money(tot['ev'])+'. Phần chi vượt giá trị đạt được là '+money(-tot['cv'])+'. AC thấp hơn PV không có nghĩa tiết kiệm, vì khối lượng hoàn thành cũng thấp hơn kế hoạch. BAC = '+money(tot['b'])+'.')
d.add_heading('Ba công việc ưu tiên xử lý',2)
p('1. G — Hồ sơ người dùng: SPI = 0,33; SV = −800; CPI = 0,80. Mới hoàn thành 1/3 dù phải xong ngày 11. G là tiền nhiệm của K và L; ưu tiên tháo gỡ G để ổn định hai nhánh checkout và thanh toán.')
p('2. F — Module đăng nhập và xác thực: SPI = 0,50; SV = −450; CPI = 0,82. Phải xong ngày 10 nhưng mới đạt 50%. F ảnh hưởng I và J, đồng thời tác động gián tiếp O qua I; cần chốt giao diện xác thực và sửa các lỗi cản trở tích hợp.')
p('3. H — Danh mục sản phẩm: SPI = 0,67; SV = −600; CPI = 0,92. Phải xong ngày 12 nhưng mới đạt 2/3. H là tiền nhiệm của M và N; hoàn tất và nghiệm thu H giúp giảm rủi ro làm lại hai nhánh này. L cần theo dõi sát vì CV = −220 và CPI = 0,80, nhưng xử lý tiền nhiệm G trước giúp giải quyết nguyên nhân ở đầu nhánh.')
d.add_heading('Nguyên nhân khả dĩ và điểm cần xác minh',2)
p('Các giả thuyết cần kiểm tra: yêu cầu hoặc API thay đổi gây làm lại; thiếu nhân lực có kinh nghiệm ở xác thực và thanh toán; lỗi tích hợp hoặc môi trường kiểm thử thiếu ổn định làm giảm năng suất. Đây là giả thuyết, chưa phải nguyên nhân đã được báo cáo xác nhận.')
p('Báo cáo ghi I/J đã làm khi F chưa xong; K/L khi G chưa xong; M/N khi H chưa xong; O khi I chưa xong. Điều này không phù hợp quan hệ FS của đề, cần xác minh việc báo cáo, nghiệm thu từng phần hoặc thay đổi cách triển khai. Giữ nguyên số liệu để tính EVM. Ngoài ra, baseline chỉ nối R với Q và nối T với S, không nối kiểm thử với các module chức năng; vì vậy đường găng 20 ngày chưa phản ánh đầy đủ điều kiện sẵn sàng phát hành thực tế.')
d.add_heading('Hành động quản lý đề xuất',2)
p('1. Trong kỳ cập nhật kế tiếp, rà soát bằng chứng hoàn thành và tiêu chí nghiệm thu với từng trưởng nhóm; thống nhất % hoàn thành, chi phí thực tế và kiểm tra các quan hệ FS bị vi phạm.')
p('2. Tập trung người có kinh nghiệm để gỡ vướng G, F, H; chia đầu việc thành hạng mục nhỏ có thể nghiệm thu, phân công chủ trì và theo dõi hằng ngày. Đánh giá chi phí bổ sung trước khi điều chuyển nhân lực.')
p('3. Chốt yêu cầu và hợp đồng API; xử lý thay đổi qua đánh giá tác động và phê duyệt. Dùng mock khi phù hợp, đồng thời duy trì review mã, kiểm thử tự động, kiểm thử hồi quy và kiểm tra bảo mật.')
p('4. Lập dự báo phần việc còn lại, rà soát điều kiện đầu vào kiểm thử/phát hành và cập nhật kế hoạch phục hồi được phê duyệt. Theo dõi CV/CPI của L, J, M và đường găng S → T; không bỏ kiểm thử để bù tiến độ, không sửa baseline gốc chỉ để che sai lệch.')
d.save(r'D:\QL dự án nhóm\output\Loi_giai_EVM_ngay_15.docx')
